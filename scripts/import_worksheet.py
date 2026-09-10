"""Send corrections made in a reading worksheet back to the record.

    python scripts/import_worksheet.py <worksheet.xlsx> --code JP-XXXX-XXXX-XXXX           # dry run
    python scripts/import_worksheet.py <worksheet.xlsx> --code JP-XXXX-XXXX-XXXX --apply   # record

Each officer the reader touched becomes a draft observation attributed to the id
code, posted through the workstation API itself - in-process, against the
database at JPOCR_DB. Vocabulary resolution, era-date parsing, ditto resolution
and every refusal are therefore exactly those of typing the officer in by hand;
nothing about what a reading means is decided here.

What is sent for an officer, field by field (standing commitment 2: a machine
value is never recorded as a person's):

  * a cell the reader changed                     - sent: it is their reading
  * a value a person had already recorded         - sent unchanged, so the new
                                                    reading (the latest wins)
                                                    does not blank what was on
                                                    record
  * a machine value the reader left untouched     - NOT sent, unless the reader
                                                    wrote "ok" in 備考, which
                                                    says "I checked this whole
                                                    row against the page"

A row nobody touched is not sent. Rows are matched by the Key column, so
sorting or filtering the sheet first does no harm; moving, inserting or renaming
columns does, and is refused.

Fields the record has no column for yet (期, 現階級任官, 前階級任官, 実役停年,
位階勲等, 生年月日) are not lost: a changed value rides in the observation's
field_confidence under "worksheet", beside the notes.
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
for sub in ("app", "reading", "ingestion"):
    sys.path.insert(0, str(ROOT / sub))

from openpyxl import load_workbook  # noqa: E402

import ditto  # noqa: E402
import page_service as ps  # noqa: E402
import worksheet  # noqa: E402

GETA = "〓"
OK_MARKS = ("ok", "ｏｋ", "✓", "✔", "確認", "確認済")
FORM_KEYS = ("seniority_no", "name_raw", "branch", "rank", "post",
             "commissioning_date", "notes")
UNSTORED = ("birth", "cohort", "rank_date", "prev_rank_date", "appointment_dates",
            "service_in_rank", "court_rank_decorations")
DITTO_COLUMN = {"branch": "branch_code", "rank": "rank_code"}


def norm(value) -> str:
    """A cell's value as the text a reader sees, for comparison and sending."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_checked(notes: str) -> bool:
    head = notes.strip().split()[0].lower() if notes.strip() else ""
    return head in OK_MARKS


@dataclass
class Plan:
    key: str
    pid: str
    frame: int
    row_index: int
    name: str
    send: dict[str, str] = field(default_factory=dict)
    why: dict[str, str] = field(default_factory=dict)
    unstored: dict[str, str] = field(default_factory=dict)
    left_as_machine: list[str] = field(default_factory=list)
    sources: dict[str, str] = field(default_factory=dict)


class WorkbookRefused(Exception):
    """The workbook cannot be matched back to what was exported."""


def read_plans(path: str | Path) -> tuple[list[Plan], list[str]]:
    """Every touched officer in the workbook, and what would be sent for each."""
    wb = load_workbook(path, data_only=True)
    try:
        reading = next(ws for ws in wb.worksheets if ws.title.startswith("読取表"))
        snapshot = wb[worksheet.SNAPSHOT]
        sources = wb[worksheet.SOURCES]
    except (StopIteration, KeyError) as exc:
        raise WorkbookRefused("not a reading worksheet, or made before corrections "
                              "could be sent back - re-export it") from exc

    keys = [c.value for c in snapshot[1]]
    cols = worksheet.columns(images="name_crop" in keys)
    if keys != [c.key for c in cols]:
        raise WorkbookRefused("this workbook came from a different version of the "
                              "exporter - re-export it")
    headers = [c.value for c in reading[1]][:len(cols)]
    if headers != [c.header for c in cols]:
        raise WorkbookRefused("columns in 読取表 have been moved, inserted or renamed - "
                              "undo that, or re-export")
    pos = {k: i + 1 for i, k in enumerate(keys)}
    editable = [c.key for c in cols if c.editable]

    plans, problems = [], []
    for r in range(2, reading.max_row + 1):
        key = reading.cell(row=r, column=pos["key"]).value
        if not key:
            continue
        try:
            origin = int(reading.cell(row=r, column=pos["row"]).value)
        except (TypeError, ValueError):
            problems.append(f"sheet row {r} ({key}): its Row# is damaged; skipped")
            continue
        if snapshot.cell(row=origin, column=pos["key"]).value != key:
            problems.append(f"sheet row {r} ({key}): does not match what was exported; "
                            f"skipped")
            continue

        now = {k: norm(reading.cell(row=r, column=pos[k]).value) for k in editable}
        was = {k: norm(snapshot.cell(row=origin, column=pos[k]).value) for k in editable}
        src = {k: sources.cell(row=origin, column=pos[k]).value or "blank"
               for k in editable}
        changed = {k for k in editable if now[k] != was[k]}
        if not changed:
            continue

        pid, frame, row_index = str(key).rsplit(":", 2)
        plan = Plan(key=key, pid=pid, frame=int(frame), row_index=int(row_index),
                    name=now.get("name_raw") or was.get("name_raw") or "", sources=src)
        checked = is_checked(now.get("notes", ""))
        for k in FORM_KEYS:
            if k in changed:
                plan.send[k], plan.why[k] = now[k], "changed"
            elif src[k] == "recorded":
                plan.send[k], plan.why[k] = now[k], "already recorded"
            elif checked and now[k]:
                plan.send[k], plan.why[k] = now[k], "row marked ok"
            elif now[k]:
                plan.left_as_machine.append(k)
        plan.unstored = {k: now[k] for k in UNSTORED if k in changed}
        plans.append(plan)
    plans.sort(key=lambda p: (p.pid, p.frame, p.row_index))
    return plans, problems


def resolve(entries: list[dict], typed: str) -> dict | None:
    """Exact label, code or recorded variant - never a near miss (see api.ts)."""
    for e in entries:
        if typed in (e.get("ja"), e.get("code")):
            return e
    for e in entries:
        if typed in (e.get("variants") or []):
            return e
    return None


def observation_body(plan: Plan, vocab: dict, source_file: str) -> dict:
    """What the workstation's form would have posted for this officer."""
    body: dict = {"row_index": plan.row_index, "ditto": []}
    confidence: dict = {}
    for k, text in plan.send.items():
        if k == "notes":
            body["notes"] = text or None
            continue
        if not text:
            continue
        if ditto.is_ditto(text):
            body["ditto"].append(DITTO_COLUMN.get(k, k))
            continue
        marks = text.count(GETA)
        if marks:
            confidence[k] = {"raw": text, "unreadable": marks}
        if k == "seniority_no":
            digits = text.translate({0xFF10 + i: ord("0") + i for i in range(10)})
            if digits.isdigit():
                body["seniority_no"] = int(digits)
            elif k not in confidence:
                confidence[k] = {"raw": text, "refused": "not a plain number"}
        elif k in ("branch", "rank"):
            kind = "branches" if k == "branch" else "ranks"
            hit = resolve(vocab.get(kind, []), text)
            if hit:
                body[f"{k}_code"] = hit["code"]
            elif k not in confidence:
                confidence[k] = {"raw": text, "refused": "not in the controlled vocabulary"}
        else:
            body[k] = text
    confidence["worksheet"] = {
        "file": source_file,
        "sent_because": plan.why,
        **({"unstored_fields": plan.unstored} if plan.unstored else {}),
    }
    body["field_confidence"] = confidence
    return body


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--code", required=True, help="your id code - the work is recorded to it")
    ap.add_argument("--apply", action="store_true",
                    help="record the changes (without it, only show what would be sent)")
    args = ap.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    try:
        plans, problems = read_plans(args.workbook)
    except WorkbookRefused as exc:
        print(f"refused: {exc}")
        return 2
    for problem in problems:
        print(f"! {problem}")
    if not plans:
        print("nothing changed in this workbook - nothing to record")
        return 0

    import warnings
    # starlette's test client warns about its httpx transport on import; it is
    # noise in a tool whose output a person reads line by line.
    warnings.filterwarnings("ignore", message=".*httpx.*")
    from fastapi.testclient import TestClient
    import api
    import db

    client = TestClient(api.app)
    headers = {"X-Annotator": args.code.strip()}
    who = client.get("/whoami", headers=headers)
    if who.status_code != 200:
        print("that id code is not recognized - nothing was sent")
        return 2
    print(f"recording as {who.json()['display_name']} into {db.db_path()}")
    print("DRY RUN - add --apply to record\n" if not args.apply else "")

    vocab = ps.vocabularies()
    name = Path(args.workbook).name
    tally: Counter = Counter()
    ensured: set[tuple[str, int]] = set()
    for plan in plans:
        body = observation_body(plan, vocab, name)
        sending = ", ".join(f"{worksheet.FIELD_JA.get(k, k)}={v or '(blank)'} [{plan.why[k]}]"
                            for k, v in plan.send.items()) or "(nothing but notes)"
        print(f"{plan.pid} frame {plan.frame} no.{plan.row_index + 1} {plan.name}")
        print(f"    send: {sending}")
        if plan.left_as_machine:
            print("    not sent (machine, untouched - write ok in 備考 to include): "
                  + "、".join(worksheet.FIELD_JA.get(k, k) for k in plan.left_as_machine))
        if plan.unstored:
            print("    kept in field_confidence (no column yet): "
                  + ", ".join(f"{worksheet.FIELD_JA.get(k, k)}={v}" for k, v in plan.unstored.items()))
        if not args.apply:
            tally["would record"] += 1
            continue

        page = (plan.pid, plan.frame)
        if page not in ensured:
            made = client.post(f"/volumes/{plan.pid}/pages/{plan.frame}/cells", headers=headers)
            if made.status_code != 201:
                print(f"    FAILED: page geometry could not be stored: {made.json().get('detail')}")
                tally["failed"] += 1
                continue
            ensured.add(page)
        saved = client.post(f"/volumes/{plan.pid}/pages/{plan.frame}/observations",
                            json=body, headers=headers)
        if saved.status_code != 201:
            print(f"    FAILED: {saved.json().get('detail')}")
            tally["failed"] += 1
            continue
        result = saved.json()
        tally["recorded"] += 1
        for column, info in (result.get("flagged") or {}).items():
            print(f"    refused {column}: {info.get('refused')}")
            tally["fields refused"] += 1
        for column, info in (result.get("inherited") or {}).items():
            print(f"    同 {column} -> {info.get('value')} (from no.{info.get('from_row', 0) + 1})")

    print("\n" + ", ".join(f"{k} {v}" for k, v in tally.items()))
    return 1 if tally.get("failed") else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

"""Tests for sending worksheet corrections back: what is sent, and what never is.

The rule under test is standing commitment 2 carried into a spreadsheet: a value
the machine read and the reader never touched must not be recorded as the
reader's, however the row around it was edited.
"""

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

import import_worksheet as I  # noqa: E402
import worksheet as W  # noqa: E402
from test_worksheet import VOCAB, sample_rows  # noqa: E402

# Sheet rows in the sample: 2 = officer 0 (a person recorded name, seniority,
# branch, rank; the machine read the date), 3 = officer 1 (machine only: a
# printed ditto), 4 = officer 2 (machine only: a refused seniority number).


class ImportPlanTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "sheet.xlsx"
        W.write_workbook(sample_rows(), [], {"pid": "p", "officers": 3}, self.path,
                         vocab=VOCAB)
        self.cols = W.columns()

    def edit(self, changes: dict):
        wb = load_workbook(self.path)
        ws = wb["読取表 Reading"]
        for (row, key), value in changes.items():
            ws.cell(row=row, column=W._index(self.cols, key), value=value)
        wb.save(self.path)

    def plans(self):
        plans, problems = I.read_plans(self.path)
        self.assertEqual(problems, [])
        return {p.key: p for p in plans}

    def test_an_untouched_workbook_sends_nothing(self):
        self.assertEqual(self.plans(), {})

    def test_a_changed_cell_is_sent_and_untouched_machine_values_are_not(self):
        self.edit({(3, "post"): "步兵第九聯隊附"})
        plan = self.plans()["p:1:1"]
        self.assertEqual(plan.send, {"post": "步兵第九聯隊附"})
        self.assertEqual(plan.why["post"], "changed")
        self.assertIn("commissioning_date", plan.left_as_machine)

    def test_values_a_person_recorded_are_carried_so_nothing_is_blanked(self):
        self.edit({(2, "post"): "步兵第七十聯隊附"})
        plan = self.plans()["p:1:0"]
        self.assertEqual(plan.why["name_raw"], "already recorded")
        self.assertEqual(plan.send["name_raw"], "平岩棟一郎")
        self.assertEqual(plan.send["branch"], "歩兵")
        self.assertNotIn("commissioning_date", plan.send)          # machine, untouched
        self.assertIn("commissioning_date", plan.left_as_machine)

    def test_ok_in_notes_sends_the_whole_checked_row(self):
        self.edit({(3, "notes"): "ok"})
        plan = self.plans()["p:1:1"]
        self.assertEqual(plan.send["commissioning_date"], "1910-12-26")
        self.assertEqual(plan.why["commissioning_date"], "row marked ok")
        self.assertEqual(plan.left_as_machine, [])

    def test_a_refusal_the_reader_corrected_is_sent(self):
        self.edit({(4, "seniority_no"): 915})
        self.assertEqual(self.plans()["p:1:2"].send["seniority_no"], "915")

    def test_a_field_with_no_record_column_is_kept_not_dropped(self):
        self.edit({(3, "cohort"): 23})
        self.assertEqual(self.plans()["p:1:1"].unstored, {"cohort": "23"})

    def test_sorting_the_sheet_first_does_no_harm(self):
        wb = load_workbook(self.path)
        ws = wb["読取表 Reading"]
        for c in range(1, ws.max_column + 1):
            top, bottom = ws.cell(row=2, column=c), ws.cell(row=4, column=c)
            top.value, bottom.value = bottom.value, top.value
        # the row now at the top is officer 2; correct its seniority there
        ws.cell(row=2, column=W._index(self.cols, "seniority_no"), value=915)
        wb.save(self.path)
        plans = self.plans()
        self.assertEqual(list(plans), ["p:1:2"])
        self.assertEqual(plans["p:1:2"].send["seniority_no"], "915")

    def test_moved_columns_are_refused_rather_than_misread(self):
        wb = load_workbook(self.path)
        wb["読取表 Reading"].insert_cols(3)
        wb.save(self.path)
        with self.assertRaises(I.WorkbookRefused):
            I.read_plans(self.path)


class ObservationBodyTests(unittest.TestCase):
    def body(self, send, **extra):
        plan = I.Plan(key="p:1:4", pid="p", frame=1, row_index=4, name="",
                      send=send, why={k: "changed" for k in send}, **extra)
        return I.observation_body(plan, VOCAB, "sheet.xlsx")

    def test_a_ditto_is_declared_not_sent_as_a_value(self):
        body = self.body({"commissioning_date": "同", "branch": "〃"})
        self.assertEqual(sorted(body["ditto"]), ["branch_code", "commissioning_date"])
        self.assertNotIn("commissioning_date", body)

    def test_branch_resolves_to_its_code_or_is_refused_with_the_reason(self):
        self.assertEqual(self.body({"branch": "歩兵"})["branch_code"], "hohei")
        refused = self.body({"branch": "騎馬"})
        self.assertNotIn("branch_code", refused)
        self.assertEqual(refused["field_confidence"]["branch"]["refused"],
                         "not in the controlled vocabulary")

    def test_full_width_seniority_digits_are_numbers(self):
        self.assertEqual(self.body({"seniority_no": "９１５"})["seniority_no"], 915)

    def test_an_unreadable_character_is_saved_with_its_count(self):
        body = self.body({"name_raw": "平〓棟一"})
        self.assertEqual(body["name_raw"], "平〓棟一")
        self.assertEqual(body["field_confidence"]["name_raw"]["unreadable"], 1)

    def test_provenance_travels_with_the_observation(self):
        body = self.body({"post": "x"}, unstored={"cohort": "23"})
        sheet = body["field_confidence"]["worksheet"]
        self.assertEqual(sheet["file"], "sheet.xlsx")
        self.assertEqual(sheet["sent_because"], {"post": "changed"})
        self.assertEqual(sheet["unstored_fields"], {"cohort": "23"})


if __name__ == "__main__":
    unittest.main()

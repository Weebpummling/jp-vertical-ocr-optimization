"""The reading worksheet: officers as a spreadsheet a person can work in.

The record lives in a SQLite file and the reading happens in a browser, but the
people this project needs can already read a roster and already know Excel.
This builds, for any run of frames of any volume, one workbook that puts
everything known about each officer on one row - what a person recorded, what
the machine proposes where nobody has yet, and where every value came from - so
it can be read, checked, sorted, shared and corrected without running any of
this code.

Sheets:

    読取表 Reading   one row per officer, in reading order. Colour says where a
                    value came from; a comment on the cell says what was read.
    凡例 Legend      the colour key, the rules, and how to send corrections back.
    ページ Pages      every frame asked for, with its completeness.
    原文 OCR          every machine reading untouched - raw text, method, reason -
                    so nothing the machine saw is lost behind a tidy cell.
    _exported        hidden: the Reading sheet exactly as written. It is how a
                    corrected cell is told apart from an untouched one, both by
                    the sheet itself (a corrected cell turns white and bold) and
                    by scripts/import_worksheet.py.
    _sources         hidden: where each of those values came from. The importer
                    needs it to keep standing commitment 2 - a machine value the
                    reader never touched is never recorded as theirs.

The colours follow the workstation's rule: a machine value must never look like
a person's. Recorded values are plain; anything the machine supplied is tinted.

A `.csv` of the Reading sheet is written beside the workbook, UTF-8 with a BOM so
Excel on a Japanese machine does not read it as cp932, with a provenance column
standing in for the colours a CSV cannot carry.
"""
from __future__ import annotations

import csv
import io
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
for sub in ("app", "reading", "ingestion"):
    if str(ROOT / sub) not in sys.path:
        sys.path.insert(0, str(ROOT / sub))

import cv2  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.properties import PageSetupProperties  # noqa: E402

import db  # noqa: E402
import page_service as ps  # noqa: E402
import proposal_service as props  # noqa: E402
import volume_service as vs  # noqa: E402

DEFAULT_BASE_URL = "http://127.0.0.1:8000"
MAX_FRAMES = 1500
MAX_IMAGE_FRAMES = 60
SNAPSHOT = "_exported"
SOURCES = "_sources"

FONT = "Yu Gothic"
ACCENT = "2F5D8A"
BODY = Font(name=FONT, size=10)
BODY_ITALIC = Font(name=FONT, size=10, italic=True)
MUTED = Font(name=FONT, size=9, color="8A8A8A")
CHECK_FONT = Font(name=FONT, size=9, color="A8541B")
LINK = Font(name=FONT, size=10, color=ACCENT, underline="single")


def _solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)


SUGGESTED = _solid("FFF6D5")     # read by machine, nobody has checked it
INHERITED = _solid("DDEBF7")     # printed as a ditto mark
REFUSED = _solid("FCE4D6")       # OCR text the rules would not accept
EDITED = _solid("FFFFFF")        # changed in this sheet - applied conditionally
HEADER = _solid(ACCENT)
FILLS = {"machine": SUGGESTED, "inherited": INHERITED, "refused": REFUSED}

PAGE_FILLS = {"complete": _solid("E2EFDA"), "leaf_missing": REFUSED,
              "in_progress": INHERITED}
PAGE_STATUS = {
    "complete": "完了 complete",
    "leaf_missing": "片丁未登録 leaf missing",
    "in_progress": "作業中 in progress",
    "not_started": "未着手 not started",
    "not_roster": "名簿外 not a roster page",
    "uncached": "画像なし not on this machine",
    "unsurveyed": "未調査 not surveyed",
}
OBS_STATUS = {
    "draft": "記録済・下書き recorded (draft)",
    "confirmed": "確認済 confirmed",
    "flagged": "要再確認 flagged",
    "adjudicated": "裁定済 adjudicated",
}

FIELD_JA = {
    "seniority_no": "序列", "name_raw": "氏名", "birth": "生年月日", "cohort": "期",
    "branch": "兵科", "rank": "階級", "post": "職名",
    "commissioning_date": "少尉任官", "rank_date": "現階級任官",
    "prev_rank_date": "前階級任官", "service_in_rank": "実役停年",
    "court_rank_decorations": "位階勲等", "notes": "備考",
}
TEMPLATE_FIELDS = ("seniority_no", "name_raw", "cohort", "post",
                   "commissioning_date", "rank_date", "prev_rank_date",
                   "service_in_rank", "court_rank_decorations")


@dataclass(frozen=True)
class Column:
    key: str
    ja: str
    en: str
    width: float
    editable: bool = False
    kind: str = "text"          # text | int
    wrap: bool = False

    @property
    def header(self) -> str:
        return f"{self.ja}\n{self.en}"


BASE_COLUMNS = (
    Column("pid", "巻", "Volume", 9),
    Column("as_of", "調", "As of", 11),
    Column("frame", "コマ", "Frame", 6),
    Column("leaf", "丁", "Leaf", 5),
    Column("officer", "番", "No.", 5),
    Column("status", "状態", "Status", 17, wrap=True),
    Column("seniority_no", "序列", "Seniority", 8, True, "int"),
    Column("name_raw", "氏名", "Name", 14, True),
    Column("birth", "生年月日", "Born (as read)", 13, True),
    Column("cohort", "期", "Cohort", 6, True, "int"),
    Column("branch", "兵科", "Branch", 8, True),
    Column("rank", "階級", "Rank", 8, True),
    Column("post", "職名", "Post", 32, True, wrap=True),
    Column("commissioning_date", "少尉任官", "Commissioned", 12, True),
    Column("rank_date", "現階級任官", "In rank since", 12, True),
    Column("prev_rank_date", "前階級任官", "Previous rank", 12, True),
    Column("service_in_rank", "実役停年", "Service in rank", 10, True),
    Column("court_rank_decorations", "位階勲等", "Court rank, orders", 18, True, wrap=True),
    Column("notes", "備考", "Notes", 22, True, wrap=True),
    Column("check", "要確認", "Check", 30, wrap=True),
    Column("image", "画像", "Image", 7),
    Column("workstation", "作業画面", "Workstation", 9),
    Column("key", "キー", "Key", 16),
    Column("row", "行番号", "Row#", 6),
)
NAME_CROP = Column("name_crop", "氏名画像", "Name (image)", 14)


def columns(images: bool = False) -> list[Column]:
    cols = list(BASE_COLUMNS)
    if images:
        at = next(i for i, c in enumerate(cols) if c.key == "name_raw") + 1
        cols.insert(at, NAME_CROP)
    return cols


def _index(cols: list[Column], key: str) -> int:
    return next(i for i, c in enumerate(cols, start=1) if c.key == key)


# --------------------------------------------------------------------------
# rows
# --------------------------------------------------------------------------

@dataclass
class OutCell:
    """One value on the sheet, and where it came from."""

    value: object = None
    source: str = "blank"       # recorded | machine | inherited | refused | blank
    comment: str | None = None
    suspect: bool = False


@dataclass
class OfficerRow:
    pid: str
    frame: int
    index: int
    panel: int
    column: int
    as_of: str | None
    status: str
    cells: dict[str, OutCell]
    checks: list[str]
    image_url: str | None = None
    workstation_url: str | None = None
    name_bbox: tuple[int, int, int, int] | None = None
    ocr: list[dict] = field(default_factory=list)

    @property
    def key(self) -> str:
        """pid:frame:row_index - the same officer the database and API name."""
        return f"{self.pid}:{self.frame}:{self.index}"


@dataclass
class PageSummary:
    frame: int
    status: str
    officers: int | None = None
    recorded: int = 0
    leaves: str = ""
    leaf_missing: bool = False
    needs_review: bool = False
    lines_outside: int | None = None
    note: str = ""


def _typed(value, kind: str):
    if kind == "int" and isinstance(value, str) and value.isdigit():
        return int(value)
    return value


def machine_cell(p: dict, kind: str = "text") -> OutCell:
    """A machine proposal as a sheet cell, tinted by how it was arrived at."""
    method = p.get("method")
    raw = (p.get("raw") or "").strip()
    value = p.get("value")
    suspect = bool(p.get("suspect"))
    if method == "blank" or (not raw and value in (None, "")):
        return OutCell(None, "blank", None, suspect)
    if method == "inherited":
        return OutCell(_typed(value, kind), "inherited",
                       f"printed as a ditto mark ({raw}): same as the officer above",
                       suspect)
    if method == "refused":
        return OutCell(raw, "refused",
                       f"OCR read: {raw}\nnot accepted: {p.get('note') or 'unreadable'}",
                       suspect)
    comment = f"OCR read: {raw}" if raw and str(value) != raw else None
    return OutCell(_typed(value, kind), "machine", comment, suspect)


def recorded_cells(obs: dict, vocab: dict | None) -> dict[str, OutCell]:
    """What a person recorded for an officer, as plain (untinted) cells."""
    vocab = vocab or {}
    branch_ja = {b["code"]: b["ja"] for b in vocab.get("branches", [])}
    rank_ja = {r["code"]: r["ja"] for r in vocab.get("ranks", [])}
    who = (f"recorded by {obs.get('author') or '(unnamed)'} at "
           f"{obs.get('created_at')} ({obs.get('status') or 'draft'})")
    out: dict[str, OutCell] = {}

    def put(key, value):
        if value not in (None, ""):
            out[key] = OutCell(value, "recorded", who)

    put("seniority_no", obs.get("seniority_no"))
    put("name_raw", obs.get("name_raw"))
    put("post", obs.get("post"))
    put("commissioning_date", obs.get("commissioning_date"))
    put("branch", branch_ja.get(obs.get("branch_code"), obs.get("branch_code")))
    put("rank", rank_ja.get(obs.get("rank_code"), obs.get("rank_code")))
    return out


def status_text(obs: dict | None, has_machine: bool) -> str:
    if obs:
        return OBS_STATUS.get(obs.get("status") or "draft",
                              f"記録済 {obs.get('status')}")
    return "機械読取のみ machine only" if has_machine else "未読 no reading"


def officer_rows(*, pid: str, frame: int, as_of: str | None,
                 page: ps.RegisteredPage, proposals: dict | None,
                 observations: list[dict], vocab: dict | None,
                 image_services: dict[int, str] | None = None,
                 base_url: str = DEFAULT_BASE_URL,
                 rereads: dict[str, dict] | None = None) -> list[OfficerRow]:
    """Every officer on one registered scan, merged from record and machine.

    A person's reading wins where there is one; the machine fills the rest, and
    each cell says which it is. Readings are append-only, so the latest one for a
    row is the one shown.
    """
    latest: dict[int, dict] = {}
    for obs in observations:
        latest[obs["row_index"]] = obs

    if proposals is None:
        unavailable = "no machine reading was attempted"
    elif not proposals.get("available", True):
        unavailable = proposals.get("reason") or "unavailable"
    else:
        unavailable = None
    by_index = {o["index"]: o for o in (proposals or {}).get("officers", [])}
    kinds = {c.key: c.kind for c in BASE_COLUMNS}
    service = (image_services or {}).get(frame)

    rows = []
    for officer in page.officers:
        prop = by_index.get(officer.index, {})
        fields = prop.get("fields", {})
        cells: dict[str, OutCell] = {}
        ocr = []
        for name in TEMPLATE_FIELDS:
            p = fields.get(name)
            if p is None:
                continue
            cells[name] = machine_cell(p, kinds.get(name, "text"))
            ocr.append({"field": name, "raw": p.get("raw"), "method": p.get("method"),
                        "value": p.get("value"), "note": p.get("note"),
                        "suspect": bool(p.get("suspect"))})
        if prop.get("birth_raw"):
            cells["birth"] = OutCell(prop["birth_raw"], "machine",
                                     "OCR read beside the name, in small type")
        has_machine = any(c.source in ("machine", "inherited", "refused")
                          for c in cells.values())

        obs = latest.get(officer.index)
        if obs:
            cells.update(recorded_cells(obs, vocab))

        checks = []
        if page.panels_missing:
            checks.append("a leaf of this scan did not register: "
                          "its officers are not in this sheet")
        shaky = [FIELD_JA.get(c.field, c.field) for c in officer.cells if c.suspect]
        if shaky:
            checks.append("cell edge inferred: " + "、".join(shaky))
        refused = [FIELD_JA.get(k, k) for k, c in cells.items() if c.source == "refused"]
        if refused:
            checks.append("OCR not accepted: " + "、".join(refused))
        differing = []
        for name in TEMPLATE_FIELDS:
            again = (rereads or {}).get(f"{officer.index}:{name}")
            if not again:
                continue
            rerun = again.get("rerun") or {}
            ocr.append({"field": name, "raw": rerun.get("raw"),
                        "method": f"ndlocr-lite/{again.get('status')}",
                        "value": rerun.get("fill"), "note": rerun.get("note"),
                        "suspect": bool(rerun.get("suspect"))})
            if again.get("status") == "alternative":
                differing.append(f"{FIELD_JA.get(name, name)} {rerun.get('fill')}")
        if differing:
            checks.append("NDLOCR-Lite, zoomed in, reads differently: " + "、".join(differing))
        if unavailable:
            checks.append("no machine reading: " + unavailable)

        x, y, w, h = officer.bbox
        rows.append(OfficerRow(
            pid=pid, frame=frame, index=officer.index, panel=officer.panel,
            column=officer.column, as_of=as_of,
            status=status_text(obs, has_machine), cells=cells, checks=checks,
            image_url=f"{service}/{x},{y},{w},{h}/full/0/default.jpg" if service else None,
            workstation_url=(f"{base_url.rstrip('/')}/?pid={pid}&frame={frame}"
                             f"&officer={officer.index + 1}"),
            name_bbox=next((tuple(c.bbox) for c in officer.cells
                            if c.field == "name_raw"), None),
            ocr=ocr,
        ))
    return rows


def page_summary(frame: int, page: ps.RegisteredPage, proposals: dict | None,
                 observations: list[dict]) -> PageSummary:
    recorded = len({o["row_index"] for o in observations})
    status = vs.page_status(vs.entry_for_page(page), recorded)["status"]
    available = bool(proposals) and proposals.get("available", True)
    if not available:
        note = "no machine reading: " + ((proposals or {}).get("reason")
                                         or "not attempted")
    else:
        note = proposals.get("ocr_note") or ""
    return PageSummary(
        frame=frame, status=status, officers=len(page.officers), recorded=recorded,
        leaves=f"{len(page.panels_registered)}/{page.panels_total}",
        leaf_missing=bool(page.panels_missing), needs_review=page.needs_review,
        lines_outside=len(proposals.get("lines_outside") or []) if available else None,
        note=note)


# --------------------------------------------------------------------------
# writing
# --------------------------------------------------------------------------

def plain_value(row: OfficerRow, key: str, excel_row: int):
    return {
        "pid": row.pid,
        "as_of": row.as_of,
        "frame": row.frame,
        "leaf": {0: "右 R", 1: "左 L"}.get(row.panel, str(row.panel)),
        "officer": row.index + 1,
        "status": row.status,
        "key": row.key,
        "row": excel_row,
    }.get(key)


def _header(ws, specs) -> None:
    for i, (text, width) in enumerate(specs, start=1):
        cell = ws.cell(row=1, column=i, value=text)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32


def _write_reading(ws, rows: list[OfficerRow], cols: list[Column],
                   name_crops: dict[str, bytes]) -> None:
    _header(ws, [(c.header, c.width) for c in cols])
    for r, row in enumerate(rows, start=2):
        tallest = 0
        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=i)
            cell.font = BODY
            cell.alignment = Alignment(vertical="top", wrap_text=col.wrap)
            if col.editable:
                out = row.cells.get(col.key)
                if out is None:
                    continue
                cell.value = out.value
                fill = FILLS.get(out.source)
                if fill is not None:
                    cell.fill = fill
                if out.suspect:
                    cell.font = BODY_ITALIC
                if out.comment:
                    note = Comment(out.comment,
                                   "record" if out.source == "recorded" else "machine")
                    note.width, note.height = 280, 110
                    cell.comment = note
            elif col.key == "image":
                if row.image_url:
                    cell.value, cell.hyperlink, cell.font = "strip", row.image_url, LINK
            elif col.key == "workstation":
                if row.workstation_url:
                    cell.value, cell.hyperlink, cell.font = ("open", row.workstation_url,
                                                             LINK)
            elif col.key == "check":
                if row.checks:
                    cell.value, cell.font = "\n".join(row.checks), CHECK_FONT
            elif col.key == "name_crop":
                png = name_crops.get(row.key)
                if png:
                    from openpyxl.drawing.image import Image as XLImage
                    image = XLImage(io.BytesIO(png))
                    image.anchor = cell.coordinate
                    ws.add_image(image)
                    tallest = max(tallest, image.height)
            else:
                cell.value = plain_value(row, col.key, r)
                if col.key in ("key", "row"):
                    cell.font = MUTED
        if tallest:
            ws.row_dimensions[r].height = tallest * 0.75 + 6

    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    ws.freeze_panes = ws.cell(row=2, column=_index(cols, "birth"))
    ws.column_dimensions[get_column_letter(_index(cols, "row"))].hidden = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"


def _write_snapshot(ws, rows: list[OfficerRow], cols: list[Column]) -> None:
    for i, col in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=col.key)
    for r, row in enumerate(rows, start=2):
        for i, col in enumerate(cols, start=1):
            if col.editable:
                out = row.cells.get(col.key)
                value = out.value if out else None
            elif col.key in ("image", "workstation", "name_crop", "check"):
                value = None
            else:
                value = plain_value(row, col.key, r)
            ws.cell(row=r, column=i, value=value)


def _write_sources(ws, rows: list[OfficerRow], cols: list[Column]) -> None:
    for i, col in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=col.key)
    for r, row in enumerate(rows, start=2):
        for i, col in enumerate(cols, start=1):
            if col.editable:
                out = row.cells.get(col.key)
                ws.cell(row=r, column=i, value=out.source if out else "blank")
            elif col.key in ("key", "row"):
                ws.cell(row=r, column=i, value=plain_value(row, col.key, r))


def _highlight_edits(ws, cols: list[Column], n: int) -> None:
    """A cell changed in the sheet turns white and bold.

    Compared against the hidden snapshot by the row's original position, held in
    the hidden Row# column, so sorting and filtering do not break it: INDEX with
    a stored row number is a direct lookup, which a MATCH on the key per cell
    would not be on a volume-sized sheet.
    """
    row_col = get_column_letter(_index(cols, "row"))
    for i, col in enumerate(cols, start=1):
        if not col.editable:
            continue
        letter = get_column_letter(i)
        rule = FormulaRule(
            formula=[f"{letter}2<>INDEX('{SNAPSHOT}'!{letter}:{letter},${row_col}2)"],
            fill=EDITED, font=Font(bold=True))
        ws.conditional_formatting.add(f"{letter}2:{letter}{n + 1}", rule)


def _vocab_dropdowns(ws, cols: list[Column], n: int, vocab: dict | None) -> None:
    """Branch and rank offer the controlled vocabulary - as a suggestion list.

    Not enforced: a printed form outside the vocabulary is a finding to flag,
    and a validation error would push a reader to bend it into something that
    fits.
    """
    if not vocab:
        return
    for key, kind in (("branch", "branches"), ("rank", "ranks")):
        labels = [e["ja"] for e in vocab.get(kind, []) if e.get("ja")]
        formula = '"' + ",".join(labels) + '"'
        if not labels or len(formula) > 255:
            continue
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showErrorMessage=False)
        letter = get_column_letter(_index(cols, key))
        dv.add(f"{letter}2:{letter}{n + 1}")
        ws.add_data_validation(dv)


def _write_pages(ws, pages: list[PageSummary]) -> None:
    _header(ws, [("コマ\nFrame", 7), ("状態\nStatus", 24), ("人数\nOfficers", 8),
                 ("記録\nRecorded", 8), ("丁\nLeaves", 7), ("片丁欠\nLeaf missing", 10),
                 ("要確認\nNeeds review", 10), ("表外OCR\nOCR outside table", 11),
                 ("備考\nNote", 70)])
    for r, page in enumerate(pages, start=2):
        values = [page.frame, PAGE_STATUS.get(page.status, page.status), page.officers,
                  page.recorded, page.leaves, "yes" if page.leaf_missing else "",
                  "yes" if page.needs_review else "", page.lines_outside, page.note]
        for i, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = BODY
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 9))
        fill = PAGE_FILLS.get(page.status)
        if fill is not None:
            ws.cell(row=r, column=2).fill = fill
    ws.freeze_panes = "B2"
    if pages:
        ws.auto_filter.ref = f"A1:I{len(pages) + 1}"


def _write_ocr(ws, rows: list[OfficerRow]) -> None:
    _header(ws, [("キー\nKey", 16), ("コマ\nFrame", 7), ("番\nNo.", 5), ("丁\nLeaf", 5),
                 ("項目\nField", 10), ("field", 22), ("原文\nRaw OCR", 26),
                 ("方法\nMethod", 10), ("候補\nProposed", 16), ("理由\nNote", 50),
                 ("辺推定\nEdge inferred", 9)])
    r = 2
    for row in rows:
        for item in row.ocr:
            values = [row.key, row.frame, row.index + 1,
                      {0: "右 R", 1: "左 L"}.get(row.panel, str(row.panel)),
                      FIELD_JA.get(item["field"], item["field"]), item["field"],
                      item["raw"], item["method"], item["value"], item["note"],
                      "yes" if item["suspect"] else ""]
            for i, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=i, value=value)
                cell.font = BODY
            r += 1
    ws.freeze_panes = "B2"
    if r > 2:
        ws.auto_filter.ref = f"A1:K{r - 1}"


def _write_legend(ws, meta: dict) -> None:
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110
    r = 1

    def line(label="", text="", *, bold=False, size=10, fill=None, font=None):
        nonlocal r
        a = ws.cell(row=r, column=1, value=label or None)
        b = ws.cell(row=r, column=2, value=text or None)
        a.font = b.font = font or Font(name=FONT, size=size, bold=bold)
        if fill is not None:
            a.fill = fill
        b.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    line("停年名簿 読取表", "Roster reading worksheet", bold=True, size=14)
    line("巻 Volume", f"{meta.get('title') or ''}  ({meta.get('pid')})")
    line("調 As of", meta.get("as_of") or "")
    line("コマ Frames", meta.get("frames_text") or "")
    line("人数 Officers", str(meta.get("officers", 0)))
    line("作成 Generated", meta.get("generated") or "")
    line()
    line("色の意味", "What the colours mean", bold=True, size=12)
    line("記録済", "Recorded by a person. Plain, exactly as in the workstation.")
    line("機械読取", "Read by machine (NDL's OCR). Nobody has checked it yet.",
         fill=SUGGESTED)
    line("同・〃", "Printed as a ditto mark: same as the officer above. The comment "
                  "shows the mark.", fill=INHERITED)
    line("不採用", "The OCR text as read, which the rules would not accept as a value "
                "(an unreadable date, letters in a number). Correct it.", fill=REFUSED)
    line("斜体 italic", "An edge of this cell was inferred rather than seen - check the "
                       "image.", font=BODY_ITALIC)
    line("白・太字", "White and bold: a cell you have changed in this sheet.",
         font=Font(name=FONT, size=10, bold=True))
    line("コメント", "Hover a cell for what was read and why; 原文 OCR has every "
                 "machine reading untouched.")
    line()
    line("作業の手順", "Working it", bold=True, size=12)
    line("1", "Open the 画像 strip link and compare each value with the print.")
    line("2", "Type over anything wrong. You do not need to change colours.")
    line("3", "〓 for a character you cannot read. 同 where the page prints a ditto.")
    line("4", "作業画面 open shows that officer in the workstation, image beside form.")
    line("5", "Never guess. A blank is better than a plausible value.")
    line("6", "Sorting and filtering are fine: rows are matched by the Key column.")
    line()
    line("記録へ戻す", "Sending corrections back to the record", bold=True, size=12)
    line("確認 dry run", "python scripts/import_worksheet.py <this file> --code <your id code>")
    line("記録 apply", "…then again with --apply. Each changed officer is recorded as a "
                     "draft under your id code, through the same rules as the "
                     "workstation.")
    line("ok", "Checked an officer against the page and found it right? Write ok in "
               "備考. Without it, values the machine read that you did not change are "
               "not recorded as yours - only what you changed is.")
    line()
    line("出典", "Provenance", bold=True, size=12)
    line("OCR", meta.get("engine") or "")
    line("画像 Images", "National Diet Library Digital Collections (IIIF). Cite the "
                       "frame URL.")


def write_workbook(rows: list[OfficerRow], pages: list[PageSummary], meta: dict,
                   path: str | Path, *, name_crops: dict[str, bytes] | None = None,
                   vocab: dict | None = None) -> Path:
    name_crops = name_crops or {}
    cols = columns(images=bool(name_crops))
    wb = Workbook()
    reading = wb.active
    reading.title = "読取表 Reading"
    _write_reading(reading, rows, cols, name_crops)
    _write_legend(wb.create_sheet("凡例 Legend"), meta)
    _write_pages(wb.create_sheet("ページ Pages"), pages)
    _write_ocr(wb.create_sheet("原文 OCR"), rows)
    snapshot = wb.create_sheet(SNAPSHOT)
    _write_snapshot(snapshot, rows, cols)
    snapshot.sheet_state = "hidden"
    sources = wb.create_sheet(SOURCES)
    _write_sources(sources, rows, cols)
    sources.sheet_state = "hidden"
    if rows:
        _highlight_edits(reading, cols, len(rows))
        _vocab_dropdowns(reading, cols, len(rows), vocab)
    wb.active = 0
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csv(rows: list[OfficerRow], path: str | Path) -> Path:
    """The Reading sheet as CSV, with provenance in place of colour."""
    cols = [c for c in BASE_COLUMNS if c.key != "row"]
    path = Path(path)
    with io.open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow([f"{c.ja} {c.en}" for c in cols] + ["出典 Sources"])
        for r, row in enumerate(rows, start=2):
            out = []
            for c in cols:
                if c.editable:
                    cell = row.cells.get(c.key)
                    out.append("" if cell is None or cell.value is None else cell.value)
                elif c.key == "check":
                    out.append(" / ".join(row.checks))
                elif c.key == "image":
                    out.append(row.image_url or "")
                elif c.key == "workstation":
                    out.append(row.workstation_url or "")
                else:
                    out.append(plain_value(row, c.key, r))
            out.append("; ".join(f"{k}={v.source}" for k, v in row.cells.items()
                                 if v.source != "blank"))
            writer.writerow(out)
    return path


# --------------------------------------------------------------------------
# the export itself
# --------------------------------------------------------------------------

@dataclass
class ExportResult:
    path: Path
    csv_path: Path
    officers: int
    frames: int
    pages: list[PageSummary]


def parse_frames(spec: str, pid: str) -> list[int]:
    """'100', '95-110', '60,95-110', or 'surveyed' (every surveyed roster page)."""
    spec = (spec or "").strip().lower()
    if not spec:
        raise ValueError("no frames given")
    if spec == "surveyed":
        survey = vs.load_survey(pid)
        frames = sorted(f for f, e in survey.items() if e.get("status") == "roster")
        if not frames:
            raise ValueError(f"no surveyed roster pages for {pid} yet - "
                             f"run: python scripts/survey_volume.py {pid}")
        return frames
    out: set[int] = set()
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            first, last = (int(x) for x in part.split("-", 1))
            if first < 1 or last < first:
                raise ValueError(f"not a frame range: {part!r}")
            out.update(range(first, last + 1))
        else:
            number = int(part)
            if number < 1:
                raise ValueError(f"not a frame: {part!r}")
            out.add(number)
    if not out:
        raise ValueError("no frames given")
    if len(out) > MAX_FRAMES:
        raise ValueError(f"{len(out)} frames asked for; the limit is {MAX_FRAMES}")
    return sorted(out)


def frames_text(frames: list[int]) -> str:
    """[60, 95, 96, 97] -> '60, 95-97'."""
    parts, start, prev = [], None, None
    for f in sorted(frames):
        if start is None:
            start = prev = f
        elif f == prev + 1:
            prev = f
        else:
            parts.append(str(start) if start == prev else f"{start}-{prev}")
            start = prev = f
    if start is not None:
        parts.append(str(start) if start == prev else f"{start}-{prev}")
    return ", ".join(parts)


def _image_services(pid: str) -> dict[int, str]:
    try:
        import iiif_client
        return {c["frame_no"]: c["service_id"]
                for c in iiif_client.canvases(iiif_client.manifest(pid))
                if c.get("service_id")}
    except (Exception, SystemExit):
        return {}


def _name_crops(image_path: Path, rows: list[OfficerRow]) -> dict[str, bytes]:
    image = cv2.imread(str(image_path), cv2.IMREAD_GRAYSCALE)
    if image is None:
        return {}
    out = {}
    for row in rows:
        if not row.name_bbox:
            continue
        x, y, w, h = row.name_bbox
        crop = image[max(y, 0):y + h, max(x, 0):x + w]
        if crop.size == 0:
            continue
        scale = min(1.0, 220 / crop.shape[0], 95 / crop.shape[1])
        if scale < 1.0:
            crop = cv2.resize(crop, None, fx=scale, fy=scale,
                              interpolation=cv2.INTER_AREA)
        ok, buf = cv2.imencode(".png", crop)
        if ok:
            out[row.key] = buf.tobytes()
    return out


def export(pid: str, frames: list[int], *, images: bool = False, fetch: bool = False,
           out_dir: str | Path | None = None, base_url: str = DEFAULT_BASE_URL,
           log=None) -> ExportResult:
    """Build the worksheet for `frames` of volume `pid`.

    Only page images already on this machine are used unless `fetch` is set.
    NDL's OCR for the volume is fetched once if missing - one request, cached
    for good - because without it there is nothing for the reader to check.
    """
    log = log or (lambda _message: None)
    if not frames:
        raise ValueError("no frames given")
    out_dir = Path(out_dir) if out_dir else vs.data_home() / "exports" / pid
    vocab = ps.vocabularies()
    volume = next((v for v in db.volumes_summary() if v["pid"] == pid), None)
    as_of = volume["edition_date"] if volume else None
    services = _image_services(pid)

    rows: list[OfficerRow] = []
    pages: list[PageSummary] = []
    crops: dict[str, bytes] = {}
    for frame in frames:
        path = vs.cache_dir(pid) / f"frame_{frame:04d}.jpg"
        if not path.exists():
            if not fetch:
                pages.append(PageSummary(frame, "uncached",
                                         note="page image not on this machine: open "
                                              "it in the workstation, or export "
                                              "with --fetch"))
                log(f"frame {frame}: not cached, skipped")
                continue
            import iiif_client
            path = iiif_client.fetch_page(pid, frame)
        try:
            page = ps.register_file(path, pid, frame)
        except ps.PageNotRegistrable as exc:
            vs.record_survey(pid, frame, vs.entry_not_roster(str(exc)))
            pages.append(PageSummary(frame, "not_roster", 0, note=str(exc)))
            log(f"frame {frame}: not a roster page")
            continue
        vs.record_page(pid, frame, page)
        try:
            proposals = props.propose_page(pid, frame, page=page)
        except props.ProposalsUnavailable as exc:
            proposals = {"available": False, "reason": str(exc), "officers": []}
        registered = db.find_page(pid, frame)
        observations = (db.observations_for_page(registered["page_id"])
                        if registered else [])
        try:
            import cell_ocr
            rereads = cell_ocr.load_results(pid, frame)
        except Exception:
            rereads = {}
        page_rows = officer_rows(pid=pid, frame=frame, as_of=as_of, page=page,
                                 proposals=proposals, observations=observations,
                                 vocab=vocab, image_services=services,
                                 base_url=base_url, rereads=rereads)
        rows.extend(page_rows)
        pages.append(page_summary(frame, page, proposals, observations))
        if images:
            crops.update(_name_crops(path, page_rows))
        log(f"frame {frame}: {len(page.officers)} officers"
            + (f", leaf missing" if page.panels_missing else ""))

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    span = f"f{frames[0]}" if len(frames) == 1 else f"f{frames[0]}-{frames[-1]}"
    path = out_dir / f"worksheet-{pid}-{span}-{stamp}.xlsx"
    meta = {
        "pid": pid,
        "title": volume["title"] if volume else "",
        "as_of": as_of,
        "frames_text": frames_text(frames),
        "officers": len(rows),
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "engine": "NDL Next Digital Library fulltext OCR (FY2021), binned into the "
                  "registered template's cells (reading/binning.py); dates read by "
                  "reading/eradate.py; ditto marks resolved per reading/ditto.py.",
    }
    write_workbook(rows, pages, meta, path, name_crops=crops, vocab=vocab)
    csv_path = write_csv(rows, path.with_suffix(".csv"))
    return ExportResult(path, csv_path, len(rows), len(frames), pages)

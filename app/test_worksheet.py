"""Tests for the reading worksheet: what lands where, in what colour, and why."""

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

import proposal_service as props
import worksheet as W
from test_proposal_service import at, spread

VOCAB = {"branches": [{"code": "hohei", "ja": "歩兵"}],
         "ranks": [{"code": "taisa", "ja": "大佐"}]}

RECORDED = {"row_index": 0, "name_raw": "平岩棟一郎", "seniority_no": 915,
            "branch_code": "hohei", "rank_code": "taisa", "post": None,
            "commissioning_date": None, "status": "draft",
            "author": "Project lead", "created_at": "2026-09-09 10:00"}


def build(lines, observations=(), *, missing_leaf=False, proposals=None):
    page = spread((2, 1))
    if missing_leaf:
        page = replace(page, panels_total=3, panels_registered=(0, 1))
    if proposals is None:
        proposals = props.propose_registered(page, lines)
        proposals["available"] = True
    rows = W.officer_rows(pid="p", frame=1, as_of="1933-09-01", page=page,
                          proposals=proposals, observations=list(observations),
                          vocab=VOCAB, image_services={1: "https://iiif.example/p/R1"},
                          base_url="http://127.0.0.1:8000")
    return page, rows


def sample_rows():
    _, rows = build([at(0, "name_raw", "平岩棟一"),
                     at(0, "seniority_no", "915"),
                     at(0, "commissioning_date", "明四三、一二、二六"),
                     at(1, "commissioning_date", "同"),
                     at(2, "seniority_no", "九一五")],
                    observations=[RECORDED])
    return rows


class RowTests(unittest.TestCase):
    def test_a_recorded_value_wins_and_says_who_recorded_it(self):
        rows = sample_rows()
        name = rows[0].cells["name_raw"]
        self.assertEqual((name.value, name.source), ("平岩棟一郎", "recorded"))
        self.assertIn("Project lead", name.comment)
        self.assertEqual(rows[0].cells["branch"].value, "歩兵")
        self.assertEqual(rows[0].cells["rank"].value, "大佐")
        self.assertIn("recorded", rows[0].status)

    def test_machine_readings_fill_only_what_nobody_recorded(self):
        rows = sample_rows()
        date = rows[0].cells["commissioning_date"]
        self.assertEqual((date.value, date.source), ("1910-12-26", "machine"))
        self.assertIn("machine only", rows[1].status)

    def test_a_ditto_is_marked_inherited_and_a_refusal_keeps_its_raw_text(self):
        rows = sample_rows()
        self.assertEqual(rows[1].cells["commissioning_date"].source, "inherited")
        self.assertEqual(rows[1].cells["commissioning_date"].value, "1910-12-26")
        refused = rows[2].cells["seniority_no"]
        self.assertEqual((refused.value, refused.source), ("九一五", "refused"))
        self.assertTrue(any("序列" in c for c in rows[2].checks))

    def test_a_missing_leaf_is_named_on_every_row(self):
        _, rows = build([], missing_leaf=True)
        for row in rows:
            self.assertTrue(any("leaf" in c for c in row.checks))

    def test_rows_carry_the_shared_key_and_links_back_to_the_image_and_workstation(self):
        rows = sample_rows()
        self.assertEqual(rows[0].key, "p:1:0")
        self.assertEqual(rows[0].image_url,
                         "https://iiif.example/p/R1/200,0,100,400/full/0/default.jpg")
        self.assertEqual(rows[2].workstation_url,
                         "http://127.0.0.1:8000/?pid=p&frame=1&officer=3")

    def test_a_zoomed_rereading_that_differs_is_named_and_every_one_is_listed(self):
        rereads = {"0:seniority_no": {"status": "alternative",
                                      "rerun": {"raw": "916", "fill": "916", "note": ""}},
                   "0:post": {"status": "agrees", "rerun": {"raw": "x", "fill": "x"}}}
        _, rows = build([at(0, "seniority_no", "915"), at(0, "post", "x")])
        page = spread((2, 1))
        proposals = props.propose_registered(page, [at(0, "seniority_no", "915")])
        rows = W.officer_rows(pid="p", frame=1, as_of=None, page=page, proposals=proposals,
                              observations=[], vocab=VOCAB, rereads=rereads)
        self.assertTrue(any("序列 916" in c for c in rows[0].checks))
        self.assertIn("ndlocr-lite/agrees", [o["method"] for o in rows[0].ocr])

    def test_a_column_marked_not_an_officer_is_left_out(self):
        page = spread((2, 1))
        proposals = props.propose_registered(page, [at(0, "seniority_no", "915")])
        rows = W.officer_rows(pid="p", frame=1, as_of=None, page=page, proposals=proposals,
                              observations=[], vocab=VOCAB, row_audit={2: "extra_row"})
        self.assertEqual([r.key for r in rows], ["p:1:0", "p:1:1"])

    def test_a_column_that_looks_like_no_officer_says_so(self):
        _, rows = build([])
        self.assertTrue(any("unused slot" in c for c in rows[0].checks))

    def test_no_machine_reading_is_said_out_loud(self):
        _, rows = build([], proposals={"available": False, "reason": "OCR not cached",
                                       "officers": []})
        self.assertTrue(any("OCR not cached" in c for c in rows[0].checks))


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.rows = sample_rows()
        self.path = Path(self.tmp.name) / "w.xlsx"
        W.write_workbook(self.rows,
                         [W.PageSummary(1, "in_progress", 3, 1, "2/2")],
                         {"pid": "p", "title": "t", "frames_text": "1",
                          "officers": len(self.rows)},
                         self.path, vocab=VOCAB)
        self.wb = load_workbook(self.path)
        self.ws = self.wb["読取表 Reading"]
        self.cols = W.columns()

    def cell(self, excel_row, key):
        return self.ws.cell(row=excel_row, column=W._index(self.cols, key))

    def test_sheets_in_order_with_the_snapshot_hidden(self):
        self.assertEqual(self.wb.sheetnames,
                         ["読取表 Reading", "凡例 Legend", "ページ Pages", "原文 OCR",
                          "_exported", "_sources"])
        self.assertEqual(self.wb["_exported"].sheet_state, "hidden")
        self.assertEqual(self.wb["_sources"].sheet_state, "hidden")

    def test_the_sources_sheet_says_where_each_value_came_from(self):
        """The importer's basis for never recording a machine value as a person's."""
        src = self.wb["_sources"]
        at_key = lambda r, key: src.cell(row=r, column=W._index(self.cols, key)).value
        self.assertEqual(at_key(2, "name_raw"), "recorded")
        self.assertEqual(at_key(2, "commissioning_date"), "machine")
        self.assertEqual(at_key(3, "commissioning_date"), "inherited")
        self.assertEqual(at_key(4, "seniority_no"), "refused")
        self.assertEqual(at_key(3, "post"), "blank")
        self.assertEqual(at_key(4, "key"), "p:1:2")

    def test_colour_says_where_a_value_came_from(self):
        self.assertIsNone(self.cell(2, "name_raw").fill.fill_type)          # recorded
        self.assertTrue(self.cell(2, "commissioning_date").fill.start_color.rgb
                        .endswith("FFF6D5"))                                 # machine
        self.assertTrue(self.cell(3, "commissioning_date").fill.start_color.rgb
                        .endswith("DDEBF7"))                                 # ditto
        self.assertTrue(self.cell(4, "seniority_no").fill.start_color.rgb
                        .endswith("FCE4D6"))                                 # refused

    def test_a_cell_comment_carries_what_was_read(self):
        self.assertIn("九一五", self.cell(4, "seniority_no").comment.text)

    def test_the_hidden_snapshot_matches_the_sheet_as_written(self):
        snap = self.wb["_exported"]
        for i, col in enumerate(self.cols, start=1):
            self.assertEqual(snap.cell(row=1, column=i).value, col.key)
            if col.editable:
                for r in range(2, len(self.rows) + 2):
                    self.assertEqual(snap.cell(row=r, column=i).value,
                                     self.ws.cell(row=r, column=i).value)

    def test_edits_are_highlighted_by_comparing_with_the_snapshot(self):
        formulas = [f for cf in self.ws.conditional_formatting
                    for rule in cf.rules for f in rule.formula]
        self.assertTrue(formulas)
        self.assertTrue(all("'_exported'!" in f for f in formulas))

    def test_branch_and_rank_suggest_the_vocabulary_without_enforcing_it(self):
        lists = self.ws.data_validations.dataValidation
        self.assertEqual(len(lists), 2)
        self.assertTrue(any("歩兵" in dv.formula1 for dv in lists))
        self.assertFalse(any(dv.showErrorMessage for dv in lists))

    def test_the_csv_has_a_bom_and_provenance_in_place_of_colour(self):
        path = W.write_csv(self.rows, Path(self.tmp.name) / "w.csv")
        raw = path.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        text = raw.decode("utf-8-sig")
        self.assertIn("出典 Sources", text)
        # officer 0: a person recorded the seniority, the machine read the date
        self.assertIn("seniority_no=recorded", text)
        self.assertIn("commissioning_date=machine", text)
        self.assertIn("seniority_no=refused", text)


class FrameSpecTests(unittest.TestCase):
    def test_single_frames_ranges_and_lists(self):
        self.assertEqual(W.parse_frames("100", "p"), [100])
        self.assertEqual(W.parse_frames("95-97", "p"), [95, 96, 97])
        self.assertEqual(W.parse_frames("60, 95-96", "p"), [60, 95, 96])

    def test_nonsense_is_refused(self):
        for spec in ("", "5-3", "0", "abc"):
            with self.subTest(spec=spec), self.assertRaises(ValueError):
                W.parse_frames(spec, "p")

    def test_frames_are_described_compactly(self):
        self.assertEqual(W.frames_text([60, 95, 96, 97]), "60, 95-97")


if __name__ == "__main__":
    unittest.main()
